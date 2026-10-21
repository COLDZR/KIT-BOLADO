"""
Sistema Financeiro Integrado — Streamlit
Unifica:
1. Leitor de Holerites
2. Dashboard Financeiro e Conciliação (Cartões)
3. Extrator de Desconto de Boletos
"""

import io
import os
import re
import sqlite3
import tempfile
import unicodedata
from contextlib import contextmanager
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl.styles.colors
import pandas as pd
import pdfplumber
import streamlit as st
from fpdf import FPDF

# ==========================================
# CONFIGURAÇÃO GERAL DA PÁGINA (Deve ser o primeiro comando)
# ==========================================
st.set_page_config(page_title="Sistema Financeiro", page_icon="💼", layout="wide")

# ==========================================
# GERAL / TRUQUES
# ==========================================
def relaxed_set(self, instance, value):
    if value is not None: pass
    instance.__dict__[self.name] = value

openpyxl.styles.colors.RGB.__set__ = relaxed_set


# ==========================================
# MÓDULO 1: LEITOR DE HOLERITES
# ==========================================
DB_PATH = Path(__file__).with_name("holerites.db")

MESES = {
    "janeiro": 1, "fevereiro": 2, "marco": 3, "abril": 4, "maio": 5, "junho": 6,
    "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12,
}

@contextmanager
def db():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = ON")
    conn.row_factory = sqlite3.Row
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()

def init_db():
    with db() as c:
        c.executescript(
            """
            CREATE TABLE IF NOT EXISTS colaboradores (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nome TEXT NOT NULL,
                cpf TEXT,
                banco TEXT,
                agencia TEXT,
                conta TEXT,
                tipo_conta TEXT DEFAULT 'Corrente'
            );
            CREATE TABLE IF NOT EXISTS holerites (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                colaborador_id INTEGER NOT NULL REFERENCES colaboradores(id) ON DELETE CASCADE,
                competencia TEXT NOT NULL,          
                liquido_centavos INTEGER NOT NULL,
                arquivo TEXT,
                criado_em TEXT,
                UNIQUE (colaborador_id, competencia)
            );
            """
        )

def listar_colaboradores() -> pd.DataFrame:
    with db() as c:
        rows = c.execute(
            "SELECT id, nome, cpf, banco, agencia, conta, tipo_conta FROM colaboradores ORDER BY nome"
        ).fetchall()
    return pd.DataFrame([dict(r) for r in rows], columns=["id", "nome", "cpf", "banco", "agencia", "conta", "tipo_conta"])

def inserir_colaborador(nome, cpf="", banco="", agencia="", conta="", tipo_conta="Corrente"):
    with db() as c:
        c.execute(
            "INSERT INTO colaboradores (nome, cpf, banco, agencia, conta, tipo_conta) VALUES (?,?,?,?,?,?)",
            (nome.strip(), somente_digitos(cpf), banco.strip(), agencia.strip(), conta.strip(), tipo_conta),
        )

def salvar_holerite(colaborador_id: int, competencia: str, liquido_centavos: int, arquivo: str):
    with db() as c:
        c.execute(
            """
            INSERT INTO holerites (colaborador_id, competencia, liquido_centavos, arquivo, criado_em)
            VALUES (?,?,?,?,?)
            ON CONFLICT(colaborador_id, competencia) DO UPDATE SET
                liquido_centavos = excluded.liquido_centavos,
                arquivo = excluded.arquivo,
                criado_em = excluded.criado_em
            """,
            (colaborador_id, competencia, liquido_centavos, arquivo, datetime.now().isoformat(timespec="seconds")),
        )

def listar_holerites() -> pd.DataFrame:
    with db() as c:
        rows = c.execute(
            """
            SELECT h.id, h.competencia, c.nome, c.cpf, c.banco, c.agencia, c.conta, c.tipo_conta,
                   h.liquido_centavos, h.arquivo, h.criado_em
            FROM holerites h JOIN colaboradores c ON c.id = h.colaborador_id
            ORDER BY h.competencia DESC, c.nome
            """
        ).fetchall()
    df = pd.DataFrame([dict(r) for r in rows])
    if not df.empty:
        df["liquido"] = df["liquido_centavos"] / 100
    return df

def normalizar(s) -> str:
    s = "" if s is None or (isinstance(s, float) and pd.isna(s)) else str(s)
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()

def somente_digitos(s) -> str:
    return re.sub(r"\D", "", "" if s is None or (isinstance(s, float) and pd.isna(s)) else str(s))

def txt(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()

def brl_para_centavos(valor: str) -> int | None:
    v = re.sub(r"[^\d,.\-]", "", valor)
    if not v:
        return None
    v = v.replace(".", "").replace(",", ".")
    try:
        return int((Decimal(v) * 100).to_integral_value())
    except InvalidOperation:
        return None

def fmt_brl(centavos) -> str:
    v = int(centavos) / 100
    return "R$ " + f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

def competencia_para_label(c: str) -> str:
    try:
        a, m = c.split("-")
        return f"{m}/{a}"
    except ValueError:
        return c

def parse_competencia_input(s: str) -> str | None:
    s = (s or "").strip()
    m = re.fullmatch(r"(0?[1-9]|1[0-2])/(\d{4})", s)
    if m:
        return f"{m.group(2)}-{int(m.group(1)):02d}"
    m = re.fullmatch(r"(\d{4})-(0[1-9]|1[0-2])", s)
    if m:
        return s
    return None

NUM = r"(\d{1,3}(?:\.\d{3})*,\d{2})"
LABELS_LIQUIDO = [
    r"l[ií]quido\s+a\s+receber", r"valor\s+l[ií]quido", r"total\s+l[ií]quido",
    r"l[ií]quido\s+a\s+pagar", r"l[ií]quido\s+do\s+m[eê]s", r"sal[aá]rio\s+l[ií]quido", r"l[ií]quido",
]

def extrair_liquido(texto: str) -> int | None:
    for label in LABELS_LIQUIDO:
        m = re.search(label + r"[^\d]{0,60}?" + NUM, texto, flags=re.IGNORECASE)
        if m:
            return brl_para_centavos(m.group(1))
    return None

def extrair_competencia(texto: str) -> str | None:
    m = re.search(r"(?:compet[êe]ncia|refer[êe]ncia|per[ií]odo)\D{0,20}(0[1-9]|1[0-2])\s*/\s*(\d{4})", texto, re.I)
    if m:
        return f"{m.group(2)}-{m.group(1)}"
    t = normalizar(texto)
    m = re.search(r"(janeiro|fevereiro|marco|abril|maio|junho|julho|agosto|setembro|outubro|novembro|dezembro)\s*(?:/|de|-)?\s*(\d{4})", t)
    if m:
        return f"{m.group(2)}-{MESES[m.group(1)]:02d}"
    m = re.search(r"\b(0[1-9]|1[0-2])/(20\d{2})\b", texto)
    if m:
        return f"{m.group(2)}-{m.group(1)}"
    return None

def identificar_colaborador(texto: str, colabs: pd.DataFrame) -> str:
    if colabs.empty:
        return ""
    digitos = somente_digitos(texto)
    for _, r in colabs.iterrows():
        cpf = somente_digitos(r["cpf"])
        if len(cpf) == 11 and cpf in digitos:
            return r["nome"]
    t = normalizar(texto)
    melhor, tam = "", 0
    for _, r in colabs.iterrows():
        n = normalizar(r["nome"])
        if n and n in t and len(n) > tam:
            melhor, tam = r["nome"], len(n)
    return melhor

def ler_pdf_holerites(arquivo, um_por_pagina: bool) -> list[tuple[str, str]]:
    nome = arquivo.name
    paginas = []
    with pdfplumber.open(io.BytesIO(arquivo.getvalue())) as pdf:
        for i, p in enumerate(pdf.pages, start=1):
            paginas.append(p.extract_text() or "")
    if um_por_pagina:
        return [(f"{nome} (pág. {i})", t) for i, t in enumerate(paginas, start=1)]
    return [(nome, "\n".join(paginas))]

CAMPOS = ["nome", "cpf", "tipo_conta", "banco", "agencia", "conta"]
CANDIDATOS = {
    "nome": ["nome", "colaborador", "funcionario", "empregado"],
    "cpf": ["cpf"],
    "tipo_conta": ["tipo de conta", "tipo conta", "tipo", "corrente/poupanca"],
    "banco": ["banco", "instituicao"],
    "agencia": ["agencia", "ag"],
    "conta": ["conta", "numero da conta", "n conta", "c/c"],
}
ROTULOS = {
    "nome": "Nome *", "cpf": "CPF", "tipo_conta": "Tipo de conta (corrente/poupança)",
    "banco": "Banco", "agencia": "Agência", "conta": "Conta",
}

def adivinhar_colunas(colunas: list[str]) -> dict:
    usadas, resultado = set(), {}
    norm = {c: normalizar(c) for c in colunas}
    for campo in CAMPOS:
        achada = None
        for cand in CANDIDATOS[campo]: 
            for c, n in norm.items():
                if c not in usadas and n == cand:
                    achada = c
                    break
            if achada: break
        if not achada:
            for cand in CANDIDATOS[campo]:
                for c, n in norm.items():
                    if c not in usadas and cand in n:
                        achada = c
                        break
                if achada: break
        resultado[campo] = achada
        if achada: usadas.add(achada)
    return resultado

def ler_planilha_holerites(arquivo) -> pd.DataFrame:
    if arquivo.name.lower().endswith((".csv", ".txt")):
        return pd.read_csv(arquivo, sep=None, engine="python", dtype=str)
    return pd.read_excel(arquivo, dtype=str)

def normalizar_tipo_conta(v) -> str:
    return "Poupança" if "poup" in normalizar(v) else "Corrente"

def importar_dataframe(df: pd.DataFrame, mapa: dict) -> tuple[int, int]:
    existentes = listar_colaboradores()
    por_cpf = {somente_digitos(r["cpf"]): r["id"] for _, r in existentes.iterrows() if somente_digitos(r["cpf"])}
    por_nome = {normalizar(r["nome"]): r["id"] for _, r in existentes.iterrows()}
    novos = atualizados = 0
    with db() as c:
        for _, linha in df.iterrows():
            def val(campo):
                col = mapa.get(campo)
                return txt(linha[col]) if col else ""
            nome = val("nome")
            if not nome: continue
            cpf = somente_digitos(val("cpf"))
            tipo = normalizar_tipo_conta(val("tipo_conta")) if mapa.get("tipo_conta") else "Corrente"
            dados = (nome, cpf, val("banco"), val("agencia"), val("conta"), tipo)
            cid = por_cpf.get(cpf) if cpf else None
            cid = cid or por_nome.get(normalizar(nome))
            if cid:
                c.execute(
                    "UPDATE colaboradores SET nome=?, cpf=?, banco=?, agencia=?, conta=?, tipo_conta=? WHERE id=?",
                    (*dados, cid),
                )
                atualizados += 1
            else:
                cur = c.execute(
                    "INSERT INTO colaboradores (nome, cpf, banco, agencia, conta, tipo_conta) VALUES (?,?,?,?,?,?)",
                    dados,
                )
                por_nome[normalizar(nome)] = cur.lastrowid
                if cpf: por_cpf[cpf] = cur.lastrowid
                novos += 1
    return novos, atualizados

def listar_ids_com_holerite(competencia: str) -> list[int]:
    with db() as c:
        return [r[0] for r in c.execute("SELECT colaborador_id FROM holerites WHERE competencia=?", (competencia,))]


# ==========================================
# MÓDULO 2: CARTÕES E CONCILIAÇÃO
# ==========================================
def formatar_moeda_cartoes(valor):
    if pd.isna(valor): return "R$ 0,00"
    return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

def limpar_valor_monetario(valor):
    if pd.isna(valor): return 0.0
    if isinstance(valor, (int, float)): return float(valor)
    v = str(valor).replace("R$", "").replace("\xa0", "").strip()
    if ',' in v and '.' in v:
        if v.rfind(',') > v.rfind('.'): v = v.replace('.', '').replace(',', '.')
        else: v = v.replace(',', '')
    elif ',' in v: v = v.replace(',', '.')
    try: return float(v)
    except: return 0.0

def encontrar_cabecalho_inteligente(df_raw):
    palavras_chave = ['data', 'valor', 'bruto', 'bandeira', 'pagamento', 'transação', 'autorização', 'terminal', 'lógico', 'nsu', 'código', 'produto']
    melhor_linha, max_matches = 0, 0
    for i in range(min(30, len(df_raw))):
        row_vals = [str(x).lower() for x in df_raw.iloc[i].values if pd.notna(x)]
        row_str = " ".join(row_vals)
        matches = sum(1 for k in palavras_chave if k in row_str)
        if matches > max_matches:
            max_matches = matches
            melhor_linha = i
    return melhor_linha

def extrair_coluna(df, possiveis_nomes):
    for col in df.columns:
        col_str = str(col).lower().strip()
        if any(p.lower() in col_str for p in possiveis_nomes):
            return col
    return None

def limpar_texto(texto):
    if pd.isna(texto): return ""
    return str(texto).encode('latin-1', 'replace').decode('latin-1')

def gerar_pdf_cartao(df_grupo, nome_grupo, qtd, bruto, mdr, liquido):
    df_grupo = df_grupo.sort_values(by='Valor bruto', ascending=True)
    pdf = FPDF(orientation='L', unit='mm', format='A4')
    pdf.add_page()
    pdf.set_font('Arial', 'B', 16)
    pdf.cell(0, 10, limpar_texto(f"Relatório de Transações - {nome_grupo}"), ln=1, align='C')
    pdf.set_font('Arial', '', 12)
    resumo = f"Lançamentos: {qtd}   |   Valor Bruto: {bruto}   |   Desconto MDR: {mdr}   |   Líquido: {liquido}"
    pdf.cell(0, 10, limpar_texto(resumo), ln=1, align='C')
    pdf.ln(5) 
    colunas = ['Pagamento', 'Venda', 'Hora', 'Autorização', 'Transação', 'Bandeira', 'V. Bruto', 'MDR', 'V. Líquido']
    larguras = [28, 25, 18, 28, 25, 25, 30, 30, 30] 
    pdf.set_font('Arial', 'B', 10)
    for i, col in enumerate(colunas):
        pdf.cell(larguras[i], 10, limpar_texto(col), border=1, align='C')
    pdf.ln()
    pdf.set_font('Arial', '', 9)
    for index, row in df_grupo.iterrows():
        pdf.cell(larguras[0], 8, limpar_texto(row['Data de pagamento']), border=1, align='C')
        pdf.cell(larguras[1], 8, limpar_texto(row['Data da venda']), border=1, align='C')
        pdf.cell(larguras[2], 8, limpar_texto(row['Hora da venda']), border=1, align='C')
        pdf.cell(larguras[3], 8, limpar_texto(row['Código de autorização']), border=1, align='C')
        pdf.cell(larguras[4], 8, limpar_texto(row['Tipo transação']), border=1, align='C')
        pdf.cell(larguras[5], 8, limpar_texto(row['Bandeira']), border=1, align='C')
        pdf.cell(larguras[6], 8, limpar_texto(formatar_moeda_cartoes(row['Valor bruto'])), border=1, align='C')
        pdf.cell(larguras[7], 8, limpar_texto(formatar_moeda_cartoes(row['Desconto MDR'])), border=1, align='C')
        pdf.cell(larguras[8], 8, limpar_texto(formatar_moeda_cartoes(row['Valor líquido'])), border=1, align='C')
        pdf.ln()
    
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        pdf.output(tmp.name)
        tmp.seek(0)
        pdf_bytes = tmp.read()
    os.unlink(tmp.name)
    return pdf_bytes

def renderizar_coluna_origem(df_origem, nome_origem):
    st.markdown(f"<h3 style='text-align: center; color: #1f77b4;'>{nome_origem}</h3>", unsafe_allow_html=True)
    if df_origem.empty:
        st.info("Nenhum dado encontrado para os filtros atuais.")
        return

    st.markdown("""
    <div style="display: flex; justify-content: space-between; border-bottom: 2px solid #ccc; padding-bottom: 5px; margin-bottom: 10px;">
        <span style="font-size: 16px;"><strong>Rótulos de Linha</strong></span>
        <span style="font-size: 16px;"><strong>Soma de Valor bruto</strong></span>
    </div>
    """, unsafe_allow_html=True)

    colunas_exibicao = ['Data da venda', 'Forma de pagamento', 'Bandeira', 'Valor bruto', 'Código de autorização', 'NSU/DOC']
    datas_unicas = sorted(df_origem['Data da venda'].dropna().unique(), reverse=True)

    for data in datas_unicas:
        df_data = df_origem[df_origem['Data da venda'] == data]
        st.markdown(f"""
        <div style="display: flex; justify-content: space-between; margin-top: 5px;">
            <span><strong>➖ {data.strftime('%d/%m/%Y')}</strong></span>
            <span><strong>{formatar_moeda_cartoes(df_data['Valor bruto'].sum())}</strong></span>
        </div>
        """, unsafe_allow_html=True)
        
        bandeiras = sorted(df_data['Bandeira'].dropna().unique())
        for bandeira in bandeiras:
            df_bandeira = df_data[df_data['Bandeira'] == bandeira]
            st.markdown(f"""
            <div style="display: flex; justify-content: space-between; margin-top: 2px;">
                <span style="margin-left: 20px;"><strong>➖ {bandeira}</strong></span>
                <span><strong>{formatar_moeda_cartoes(df_bandeira['Valor bruto'].sum())}</strong></span>
            </div>
            """, unsafe_allow_html=True)
            
            formas = sorted(df_bandeira['Forma de pagamento'].dropna().unique())
            for forma in formas:
                df_forma = df_bandeira[df_bandeira['Forma de pagamento'] == forma]
                total_forma = df_forma['Valor bruto'].sum()
                with st.expander(f"    {forma} ⸻ {formatar_moeda_cartoes(total_forma)}"):
                    st.dataframe(
                        df_forma[colunas_exibicao], 
                        use_container_width=True, hide_index=True,
                        column_config={
                            "Data da venda": st.column_config.DateColumn("Data", format="DD/MM/YYYY"),
                            "Valor bruto": st.column_config.NumberColumn("Valor bruto", format="R$ %.2f")
                        }
                    )
        st.markdown("<hr style='margin: 10px 0; border: 0; border-top: 1px dashed #eee;'>", unsafe_allow_html=True)


# ==========================================
# MÓDULO 3: DESCONTOS DE BOLETOS
# ==========================================
LINHA_CABECALHO_RE = re.compile(r"(\d{2}/\d{2}/\d{4})\s+(\S+)\s+(NF|ND|DM|DS)\s+([SN])\s+(\d{2}/\d{2}/\d{4})\s+(\S+)")
VALOR_DOCUMENTO_RE = re.compile(r"Valor do Documento\n(.+)")
VALOR_ATE_VENCIMENTO_RE = re.compile(r"Até o vencimento cobrar:\s*R\$\s*([\d\.]+,\d{2})")
MONEY_RE = re.compile(r"([\d\.]+,\d{2})")

def valor_br_para_float(valor_str: str) -> float:
    return float(valor_str.replace(".", "").replace(",", "."))

def extrair_boletos(pdf_bytes: bytes) -> pd.DataFrame:
    texto_completo = ""
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for pagina in pdf.pages:
            texto_pagina = pagina.extract_text() or ""
            texto_completo += texto_pagina + "\n"
    
    blocos = re.split(r"(?=Local de Vencimento\n)", texto_completo)
    registros = {} 
    
    for bloco in blocos:
        m_cabecalho = LINHA_CABECALHO_RE.search(bloco)
        if not m_cabecalho: continue
        data_documento, numero_documento, nosso_numero = m_cabecalho.group(1), m_cabecalho.group(2), m_cabecalho.group(6)
        
        m_valor_doc = VALOR_DOCUMENTO_RE.search(bloco)
        if not m_valor_doc: continue
        
        valores_encontrados = MONEY_RE.findall(m_valor_doc.group(1))
        if not valores_encontrados: continue
        valor_documento = valor_br_para_float(valores_encontrados[-1])
        
        m_valor_vencimento = VALOR_ATE_VENCIMENTO_RE.search(bloco)
        valor_ate_vencimento = valor_br_para_float(m_valor_vencimento.group(1)) if m_valor_vencimento else None
        
        desconto = round(valor_documento - valor_ate_vencimento, 2) if valor_ate_vencimento is not None else 0.0
        
        registros[nosso_numero] = {
            "Nosso Número": nosso_numero, "Número do Documento": numero_documento,
            "Data do Documento": data_documento, "Valor do Documento (R$)": valor_documento,
            "Valor Até o Vencimento (R$)": valor_ate_vencimento, "Desconto (R$)": desconto,
        }
    return pd.DataFrame(list(registros.values()))

def gerar_excel_boletos(df: pd.DataFrame) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Descontos")
        worksheet = writer.sheets["Descontos"]
        for i, col in enumerate(df.columns, start=1):
            largura_max_valores = df[col].apply(lambda x: len(str(x))).max()
            worksheet.column_dimensions[worksheet.cell(row=1, column=i).column_letter].width = max(largura_max_valores, len(col)) + 2
    buffer.seek(0)
    return buffer.getvalue()


# ==========================================
# INTERFACE PRINCIPAL DO APP
# ==========================================
def main():
    init_db()
    
    st.title("💼 Sistema Financeiro Integrado")
    st.markdown("Selecione a ferramenta desejada nas abas abaixo.")
    
    # Abas principais unindo os 3 apps
    tab_holerites, tab_cartoes, tab_boletos = st.tabs([
        "🧾 Leitor de Holerites", 
        "📈 Dashboard Financeiro e Cartões", 
        "🎫 Extrator de Boletos"
    ])
    
    # ---------------------------------------------------------
    # ABA PRINCIPAL 1: HOLERITES
    # ---------------------------------------------------------
    with tab_holerites:
        st.header("Leitor de Holerites")
        st.caption("Extrai o valor líquido de cada colaborador a partir dos PDFs e guarda por competência.")
        t1, t2, t3 = st.tabs(["👥 Colaboradores", "📄 Ler holerites", "💰 Resultados"])
        
        # ABA: COLABORADORES
        with t1:
            st.subheader("Cadastro de colaboradores")
            modo = st.radio("Como deseja cadastrar?", ["Cadastrar manualmente", "Importar de planilha existente"], horizontal=True)
            if modo == "Cadastrar manualmente":
                with st.form("form_colab", clear_on_submit=True):
                    c1, c2 = st.columns(2)
                    nome, cpf = c1.text_input("Nome completo *"), c2.text_input("CPF")
                    c3, c4, c5, c6 = st.columns([2, 1, 1.5, 1.5])
                    banco, agencia, conta = c3.text_input("Banco"), c4.text_input("Agência"), c5.text_input("Conta")
                    tipo = c6.selectbox("Tipo de conta", ["Corrente", "Poupança"])
                    if st.form_submit_button("Cadastrar", type="primary"):
                        if not nome.strip(): st.error("Informe o nome do colaborador.")
                        else:
                            inserir_colaborador(nome, cpf, banco, agencia, conta, tipo)
                            st.success(f"{nome.strip()} cadastrado(a).")
                            st.rerun()
            else:
                st.caption("Envie uma planilha (.xlsx, .xls ou .csv) com os colaboradores e seus dados bancários.")
                arq = st.file_uploader("Planilha", type=["xlsx", "xls", "csv"], key="up_planilha")
                if arq:
                    try: df_colabs_plan = ler_planilha_holerites(arq)
                    except Exception as e: st.error(f"Não consegui ler a planilha: {e}")
                    else:
                        st.write(f"{len(df_colabs_plan)} linhas encontradas. Confira o mapeamento das colunas:")
                        colunas = list(df_colabs_plan.columns)
                        palpite = adivinhar_colunas(colunas)
                        opcoes, mapa = ["— não usar —"] + colunas, {}
                        cols = st.columns(3)
                        for i, campo in enumerate(CAMPOS):
                            idx = opcoes.index(palpite[campo]) if palpite[campo] in opcoes else 0
                            esc = cols[i % 3].selectbox(ROTULOS[campo], opcoes, index=idx, key=f"map_{campo}")
                            mapa[campo] = None if esc == opcoes[0] else esc
                        st.dataframe(df_colabs_plan.head(10), use_container_width=True)
                        if st.button("Importar colaboradores", type="primary"):
                            if not mapa["nome"]: st.error("Selecione a coluna do nome.")
                            else:
                                novos, atual = importar_dataframe(df_colabs_plan, mapa)
                                st.success(f"Importação concluída: {novos} novo(s), {atual} atualizado(s).")
                                st.rerun()

            st.divider()
            st.subheader("Colaboradores cadastrados")
            colabs = listar_colaboradores()
            if colabs.empty: st.info("Nenhum colaborador cadastrado ainda.")
            else:
                colabs["excluir"] = False
                editado = st.data_editor(
                    colabs, hide_index=True, use_container_width=True, disabled=["id"],
                    column_config={
                        "nome": "Nome", "cpf": "CPF", "banco": "Banco", "agencia": "Agência", "conta": "Conta",
                        "tipo_conta": st.column_config.SelectboxColumn("Tipo", options=["Corrente", "Poupança"]),
                        "excluir": st.column_config.CheckboxColumn("Excluir"),
                    },
                    key="editor_colabs",
                )
                if st.button("Salvar alterações"):
                    with db() as c:
                        for _, r in editado.iterrows():
                            if r["excluir"]: c.execute("DELETE FROM colaboradores WHERE id=?", (int(r["id"]),))
                            else:
                                c.execute("UPDATE colaboradores SET nome=?, cpf=?, banco=?, agencia=?, conta=?, tipo_conta=? WHERE id=?",
                                          (txt(r["nome"]), somente_digitos(r["cpf"]), txt(r["banco"]), txt(r["agencia"]),
                                           txt(r["conta"]), txt(r["tipo_conta"]) or "Corrente", int(r["id"])))
                    st.success("Alterações salvas.")
                    st.rerun()

        # ABA: LER HOLERITES
        with t2:
            st.subheader("Ler holerites")
            colabs = listar_colaboradores()
            if colabs.empty: st.warning("Cadastre ou importe os colaboradores primeiro (aba Colaboradores).")
            else:
                c1, c2 = st.columns([2, 1])
                arquivos = c1.file_uploader("PDFs dos holerites", type=["pdf"], accept_multiple_files=True)
                um_por_pagina = c2.checkbox("PDF com vários holerites (um por página)", value=True)
                comp_padrao = c2.text_input("Competência padrão (MM/AAAA)", placeholder="usada se não achar no PDF")

                if arquivos and st.button("Ler holerites", type="primary"):
                    linhas, avisos = [], []
                    for arq in arquivos:
                        try: unidades = ler_pdf_holerites(arq, um_por_pagina)
                        except Exception as e:
                            avisos.append(f"{arq.name}: erro ao abrir o PDF ({e})")
                            continue
                        for rotulo, texto in unidades:
                            if not texto.strip(): avisos.append(f"{rotulo}: sem texto (PDF escaneado?). Preencha manualmente ou use um PDF com texto.")
                            comp = extrair_competencia(texto) or parse_competencia_input(comp_padrao)
                            liq = extrair_liquido(texto)
                            linhas.append({
                                "arquivo": rotulo, "colaborador": identificar_colaborador(texto, colabs),
                                "competencia": competencia_para_label(comp) if comp else "",
                                "liquido": (liq / 100) if liq is not None else None,
                            })
                    st.session_state["lidos"] = pd.DataFrame(linhas)
                    for a in avisos: st.warning(a)

                lidos = st.session_state.get("lidos")
                if lidos is not None and not lidos.empty:
                    st.write("Revise os dados lidos. Corrija o que for preciso antes de salvar:")
                    nomes = [""] + colabs["nome"].tolist()
                    editado = st.data_editor(
                        lidos, hide_index=True, use_container_width=True, num_rows="dynamic",
                        column_config={
                            "arquivo": st.column_config.TextColumn("Arquivo", disabled=True),
                            "colaborador": st.column_config.SelectboxColumn("Colaborador", options=nomes),
                            "competencia": st.column_config.TextColumn("Competência (MM/AAAA)"),
                            "liquido": st.column_config.NumberColumn("Líquido (R$)", format="R$ %.2f", min_value=0.0, step=0.01),
                        },
                        key="editor_lidos",
                    )
                    ids = {r["nome"]: int(r["id"]) for _, r in colabs.iterrows()}
                    if st.button("Salvar valores líquidos", type="primary"):
                        salvos, ignorados = 0, []
                        for _, r in editado.iterrows():
                            comp = parse_competencia_input(txt(r.get("competencia")))
                            liq, nome = r.get("liquido"), txt(r.get("colaborador"))
                            if not nome or nome not in ids or comp is None or liq is None or pd.isna(liq):
                                ignorados.append(txt(r.get("arquivo")) or "(linha sem arquivo)")
                                continue
                            salvar_holerite(ids[nome], comp, int(round(float(liq) * 100)), txt(r.get("arquivo")))
                            salvos += 1
                        st.success(f"{salvos} holerite(s) salvo(s).")
                        if ignorados: st.warning("Ignorados por falta de colaborador, competência ou valor: " + ", ".join(ignorados))
                        else:
                            st.session_state.pop("lidos", None)
                            st.rerun()

        # ABA: RESULTADOS
        with t3:
            st.subheader("Valores líquidos armazenados")
            df_resultados = listar_holerites()
            if df_resultados.empty: st.info("Nenhum holerite salvo ainda.")
            else:
                comps = sorted(df_resultados["competencia"].unique(), reverse=True)
                escolha = st.selectbox("Competência", ["Todas"] + comps, format_func=lambda c: c if c == "Todas" else competencia_para_label(c))
                vis = df_resultados if escolha == "Todas" else df_resultados[df_resultados["competencia"] == escolha]

                m1, m2 = st.columns(2)
                m1.metric("Holerites", len(vis))
                m2.metric("Total líquido", fmt_brl(vis["liquido_centavos"].sum()))

                tabela = vis[["competencia", "nome", "cpf", "banco", "agencia", "conta", "tipo_conta", "liquido"]].copy()
                tabela["competencia"] = tabela["competencia"].map(competencia_para_label)
                tabela.columns = ["Competência", "Colaborador", "CPF", "Banco", "Agência", "Conta", "Tipo", "Líquido"]
                st.dataframe(tabela, hide_index=True, use_container_width=True, column_config={"Líquido": st.column_config.NumberColumn(format="R$ %.2f")})

                buf = io.BytesIO()
                with pd.ExcelWriter(buf, engine="openpyxl") as w:
                    tabela.to_excel(w, index=False, sheet_name="Líquidos")
                d1, d2 = st.columns(2)
                d1.download_button("Baixar Excel", buf.getvalue(), "liquidos.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                d2.download_button("Baixar CSV", tabela.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig"), "liquidos.csv", "text/csv")

                if escolha != "Todas":
                    colabs = listar_colaboradores()
                    faltam = colabs[~colabs["id"].isin([int(i) for i in listar_ids_com_holerite(escolha)])]
                    if not faltam.empty:
                        with st.expander(f"Sem holerite em {competencia_para_label(escolha)} ({len(faltam)})"):
                            st.write(", ".join(faltam["nome"]))

                with st.expander("Excluir um registro"):
                    opcoes_excl = {f"{competencia_para_label(r['competencia'])} — {r['nome']} — {fmt_brl(r['liquido_centavos'])}": int(r["id"]) for _, r in df_resultados.iterrows()}
                    sel = st.selectbox("Registro", list(opcoes_excl))
                    if st.button("Excluir registro"):
                        with db() as c: c.execute("DELETE FROM holerites WHERE id=?", (opcoes_excl[sel],))
                        st.rerun()


    # ---------------------------------------------------------
    # ABA PRINCIPAL 2: CARTÕES E CONCILIAÇÃO
    # ---------------------------------------------------------
    with tab_cartoes:
        st.header("Dashboard Financeiro e Conciliação")
        tab_c1, tab_c2 = st.tabs(["📊 Resumo de Recebimentos (Sicredi)", "🔄 Conciliação de Vendas (Tabela Dinâmica)"])
        
        # ABA: RESUMO SICREDI
        with tab_c1:
            st.markdown("Faça o upload do seu relatório em Excel (.xlsx) para visualizar a análise detalhada e gerar os PDFs individuais.")
            uploaded_file = st.file_uploader("Selecione a planilha Sicredi (.xlsx)", type=["xlsx"], key="up_sicredi_aba1")

            if uploaded_file is not None:
                try:
                    df_raw = pd.read_excel(uploaded_file, header=None)
                    header_row = encontrar_cabecalho_inteligente(df_raw)
                    df_full = df_raw.iloc[header_row + 1:].copy()
                    df_full.columns = df_raw.iloc[header_row].astype(str).str.strip()
                    
                    target_cols = ['Data de pagamento', 'Data da venda', 'Hora da venda', 'Código de autorização', 'Tipo transação', 'Bandeira', 'Valor bruto', 'Desconto MDR', 'Valor líquido']
                    colunas_faltantes = [col for col in target_cols if col not in df_full.columns.tolist()]
                    
                    if colunas_faltantes:
                        st.error(f"⚠️ Colunas não encontradas: {', '.join(colunas_faltantes)}")
                    else:
                        df_filtered = df_full[target_cols].copy()
                        for col in ['Valor bruto', 'Desconto MDR', 'Valor líquido']:
                            df_filtered[col] = pd.to_numeric(df_filtered[col].astype(str).str.replace(',', '.'), errors='coerce')
                        
                        df_filtered = df_filtered.dropna(subset=['Bandeira', 'Tipo transação'])
                        df_filtered['Combinação'] = df_filtered['Bandeira'].astype(str).str.upper() + " " + df_filtered['Tipo transação'].astype(str).str.upper()
                        df_filtered['Filtro Data'] = pd.to_datetime(df_filtered['Data de pagamento'], dayfirst=True, errors='coerce')
                        datas_validas = df_filtered['Filtro Data'].dropna()
                        
                        if not datas_validas.empty:
                            min_date, max_date = datas_validas.min().date(), datas_validas.max().date()
                            st.sidebar.header("Filtros - Recebimentos (Cartões)")
                            data_selecionada = st.sidebar.date_input("Data de Pagamento", value=(min_date, max_date), min_value=min_date, max_value=max_date)
                            
                            if len(data_selecionada) == 2:
                                mask = (df_filtered['Filtro Data'].dt.date >= data_selecionada[0]) & (df_filtered['Filtro Data'].dt.date <= data_selecionada[1])
                                df_filtered = df_filtered.loc[mask]
                        
                        df_filtered = df_filtered.drop(columns=['Filtro Data'])

                        col1_1, col2_1, col3_1, col4_1 = st.columns(4)
                        col1_1.metric("Lançamentos", f"{len(df_filtered)}")
                        col2_1.metric("Bruto", formatar_moeda_cartoes(df_filtered['Valor bruto'].sum()))
                        col3_1.metric("Desconto MDR", formatar_moeda_cartoes(df_filtered['Desconto MDR'].sum()))
                        col4_1.metric("Líquido", formatar_moeda_cartoes(df_filtered['Valor líquido'].sum()))
                        
                        st.divider()
                        st.subheader("Visualização em Planilha")
                        st.dataframe(df_filtered.drop(columns=['Combinação']), use_container_width=True, hide_index=True)
                        
                        st.divider()
                        st.subheader("Resumo por Combinação (Cartões)")
                        if df_filtered.empty: st.warning("Não há dados para o período selecionado.")
                        else:
                            grupos = df_filtered.groupby('Combinação')
                            cols_cards = st.columns(3)
                            col_idx = 0
                            for nome_grupo, df_grupo in grupos:
                                qtd = len(df_grupo)
                                str_bruto = formatar_moeda_cartoes(df_grupo['Valor bruto'].sum())
                                str_mdr = formatar_moeda_cartoes(df_grupo['Desconto MDR'].sum())
                                str_liquido = formatar_moeda_cartoes(df_grupo['Valor líquido'].sum())
                                
                                with cols_cards[col_idx % 3]:
                                    st.markdown(f"""
                                    <div style="border: 1px solid #e6e6e6; border-radius: 10px; padding: 20px; margin-bottom: 15px; box-shadow: 2px 2px 5px rgba(0,0,0,0.05);">
                                        <h4 style="color: #2c3e50; margin-top: 0; border-bottom: 2px solid #f0f2f6; padding-bottom: 10px;">💳 {nome_grupo}</h4>
                                        <p style="margin: 8px 0;"><strong>Lançamentos:</strong> {qtd}</p>
                                        <p style="margin: 8px 0; color: #555;"><strong>Bruto:</strong> {str_bruto}</p>
                                        <p style="margin: 8px 0; color: #d9534f;"><strong>MDR:</strong> {str_mdr}</p>
                                        <p style="margin: 8px 0; color: #27ae60;"><strong>Líquido:</strong> {str_liquido}</p>
                                    </div>
                                    """, unsafe_allow_html=True)
                                    
                                    pdf_bytes = gerar_pdf_cartao(df_grupo.drop(columns=['Combinação']), nome_grupo, qtd, str_bruto, str_mdr, str_liquido)
                                    st.download_button(label=f"📄 Baixar PDF ({nome_grupo})", data=pdf_bytes, file_name=f"Relatorio_{nome_grupo.replace(' ', '_')}.pdf", mime='application/pdf', key=f"btn_{nome_grupo}", use_container_width=True)
                                col_idx += 1
                except Exception as e: st.error(f"Erro ao processar arquivo: {e}")

        # ABA: CONCILIAÇÃO
        with tab_c2:
            st.markdown("Faça o upload dos 3 relatórios **em formato .xlsx** para conciliar as vendas.")
            col_up1, col_up2, col_up3 = st.columns(3)
            with col_up1: file_cielo = st.file_uploader("1. Relatório Cielo (.xlsx)", type=["xlsx"], key="up_cielo")
            with col_up2: file_sicredi = st.file_uploader("2. Relatório Sicredi (.xlsx)", type=["xlsx"], key="up_sicredi")
            with col_up3: file_pix = st.file_uploader("3. Relatório Sicredi PIX (.xlsx)", type=["xlsx"], key="up_pix")

            if file_cielo and file_sicredi and file_pix:
                with st.spinner('Lendo os arquivos Excel de forma robusta...'):
                    try:
                        colunas_padrao = ['Data da venda', 'Forma de pagamento', 'Bandeira', 'Valor bruto', 'Código de autorização', 'NSU/DOC', 'Máquina', 'Origem']
                        dataframes_processados = []

                        # --- CIELO ---
                        df_raw_cielo = pd.read_excel(file_cielo, header=None)
                        h_cielo = encontrar_cabecalho_inteligente(df_raw_cielo)
                        df_cielo = df_raw_cielo.iloc[h_cielo + 1:].copy()
                        df_cielo.columns = df_raw_cielo.iloc[h_cielo].astype(str).str.strip()
                        df_c_padrao = pd.DataFrame(columns=colunas_padrao)
                        col_data = extrair_coluna(df_cielo, ['data da venda', 'data'])
                        col_forma = extrair_coluna(df_cielo, ['forma de pagamento', 'produto'])
                        col_band = extrair_coluna(df_cielo, ['bandeira'])
                        col_val = extrair_coluna(df_cielo, ['valor bruto', 'bruto'])
                        col_aut = extrair_coluna(df_cielo, ['autorização', 'cód. autorização'])
                        col_nsu = extrair_coluna(df_cielo, ['nsu', 'doc'])
                        col_maq = extrair_coluna(df_cielo, ['lógico', 'nº lógico', 'terminal', 'equipamento'])
                        if col_data: df_c_padrao['Data da venda'] = pd.to_datetime(df_cielo[col_data], dayfirst=True, errors='coerce').dt.date
                        if col_forma: df_c_padrao['Forma de pagamento'] = df_cielo[col_forma].astype(str).str.strip()
                        if col_band: df_c_padrao['Bandeira'] = df_cielo[col_band]
                        if col_val: df_c_padrao['Valor bruto'] = df_cielo[col_val].apply(limpar_valor_monetario)
                        if col_aut: df_c_padrao['Código de autorização'] = df_cielo[col_aut]
                        if col_nsu: df_c_padrao['NSU/DOC'] = df_cielo[col_nsu]
                        if col_maq: df_c_padrao['Máquina'] = df_cielo[col_maq]
                        df_c_padrao['Origem'] = 'Cielo'
                        df_c_padrao['Forma de pagamento'] = df_c_padrao['Forma de pagamento'].replace({'Débito pré-pago': 'Débito à vista', 'Crédito pré-pago': 'Crédito à vista'})
                        dataframes_processados.append(df_c_padrao.dropna(subset=['Valor bruto']))

                        # --- SICREDI CARTÕES ---
                        df_raw_sicredi = pd.read_excel(file_sicredi, header=None)
                        h_sicredi = encontrar_cabecalho_inteligente(df_raw_sicredi) 
                        df_sicredi = df_raw_sicredi.iloc[h_sicredi + 1:].copy()
                        df_sicredi.columns = df_raw_sicredi.iloc[h_sicredi].astype(str).str.strip()
                        df_s_padrao = pd.DataFrame(columns=colunas_padrao)
                        col_data_s = extrair_coluna(df_sicredi, ['data da venda'])
                        col_forma_s = extrair_coluna(df_sicredi, ['tipo transação'])
                        col_band_s = extrair_coluna(df_sicredi, ['bandeira'])
                        col_val_s = extrair_coluna(df_sicredi, ['valor bruto'])
                        col_aut_s = extrair_coluna(df_sicredi, ['código de autorização', 'autorização'])
                        col_nsu_s = extrair_coluna(df_sicredi, ['comprovante da venda', 'nsu'])
                        col_maq_s = extrair_coluna(df_sicredi, ['número do terminal', 'terminal'])
                        if col_data_s: df_s_padrao['Data da venda'] = pd.to_datetime(df_sicredi[col_data_s], dayfirst=True, errors='coerce').dt.date
                        if col_forma_s: df_s_padrao['Forma de pagamento'] = df_sicredi[col_forma_s]
                        if col_band_s: df_s_padrao['Bandeira'] = df_sicredi[col_band_s]
                        if col_val_s: df_s_padrao['Valor bruto'] = df_sicredi[col_val_s].apply(limpar_valor_monetario)
                        if col_aut_s: df_s_padrao['Código de autorização'] = df_sicredi[col_aut_s]
                        if col_nsu_s: df_s_padrao['NSU/DOC'] = df_sicredi[col_nsu_s]
                        if col_maq_s: df_s_padrao['Máquina'] = df_sicredi[col_maq_s]
                        df_s_padrao['Origem'] = 'Sicredi'
                        dataframes_processados.append(df_s_padrao.dropna(subset=['Valor bruto']))

                        # --- SICREDI PIX ---
                        df_raw_pix = pd.read_excel(file_pix, header=None)
                        h_pix = encontrar_cabecalho_inteligente(df_raw_pix)
                        df_pix = df_raw_pix.iloc[h_pix + 1:].copy()
                        df_pix.columns = df_raw_pix.iloc[h_pix].astype(str).str.strip()
                        df_p_padrao = pd.DataFrame(columns=colunas_padrao)
                        col_data_p = extrair_coluna(df_pix, ['data da venda', 'data'])
                        col_val_p = extrair_coluna(df_pix, ['valor bruto', 'bruto'])
                        col_aut_p = extrair_coluna(df_pix, ['autorização', 'cód'])
                        col_maq_p = extrair_coluna(df_pix, ['terminal'])
                        if col_data_p: df_p_padrao['Data da venda'] = pd.to_datetime(df_pix[col_data_p], dayfirst=True, errors='coerce').dt.date
                        df_p_padrao['Forma de pagamento'] = 'PIX'
                        df_p_padrao['Bandeira'] = 'PIX'
                        if col_val_p: df_p_padrao['Valor bruto'] = df_pix[col_val_p].apply(limpar_valor_monetario)
                        if col_aut_p: 
                            df_p_padrao['Código de autorização'] = df_pix[col_aut_p]
                            df_p_padrao['NSU/DOC'] = df_pix[col_aut_p] 
                        if col_maq_p: df_p_padrao['Máquina'] = df_pix[col_maq_p]
                        df_p_padrao['Origem'] = 'Sicredi PIX'
                        dataframes_processados.append(df_p_padrao.dropna(subset=['Valor bruto']))

                        df_final = pd.concat(dataframes_processados, ignore_index=True)
                        df_final['Máquina'] = df_final['Máquina'].fillna('Não identificada').astype(str).str.replace(r'\.0$', '', regex=True)
                        
                        st.divider()
                        st.sidebar.header("Filtros - Conciliação (Cartões)")
                        maquinas_disponiveis = sorted(df_final['Máquina'].unique().tolist())
                        maquinas_selecionadas = st.sidebar.multiselect("Selecione os Terminais / Máquinas", options=maquinas_disponiveis, default=maquinas_disponiveis)
                        
                        if maquinas_selecionadas: df_final = df_final[df_final['Máquina'].isin(maquinas_selecionadas)]
                        
                        col_c, col_s, col_p = st.columns(3)
                        with col_c: renderizar_coluna_origem(df_final[df_final['Origem'] == 'Cielo'], "🔵 CIELO")
                        with col_s: renderizar_coluna_origem(df_final[df_final['Origem'] == 'Sicredi'], "🟢 SICREDI CARTÕES")
                        with col_p: renderizar_coluna_origem(df_final[df_final['Origem'] == 'Sicredi PIX'], "💠 SICREDI PIX")

                    except Exception as e: st.error(f"Erro ao consolidar planilhas. Detalhe: {e}")

    # ---------------------------------------------------------
    # ABA PRINCIPAL 3: DESCONTOS DE BOLETOS
    # ---------------------------------------------------------
    with tab_boletos:
        st.header("Extrator de Desconto de Boletos")
        st.caption(
            "Envie um PDF com vários boletos (ex: lote de extrato de boletos do Bradesco/Banco do Brasil). "
            "O app identifica cada boleto único e calcula o desconto: **Valor do Documento − Valor Até o Vencimento**."
        )

        arquivo_pdf = st.file_uploader("Selecione o arquivo PDF", type=["pdf"], key="up_boletos")

        if arquivo_pdf is not None:
            with st.spinner("Lendo o PDF e extraindo os boletos..."):
                try: df_boletos = extrair_boletos(arquivo_pdf.read())
                except Exception as e:
                    st.error(f"Ocorreu um erro ao processar o PDF: {e}")
                    df_boletos = pd.DataFrame()

            if df_boletos.empty:
                st.warning("Nenhum boleto foi identificado neste PDF. Verifique se o layout do arquivo corresponde ao esperado.")
            else:
                total_documentos = df_boletos["Valor do Documento (R$)"].sum()
                total_desconto = df_boletos["Desconto (R$)"].sum()

                col1_bol, col2_bol, col3_bol = st.columns(3)
                col1_bol.metric("Boletos encontrados", len(df_boletos))
                col2_bol.metric("Total Valor Documento", f"R$ {total_documentos:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
                col3_bol.metric("Total Desconto", f"R$ {total_desconto:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))

                st.dataframe(
                    df_boletos.style.format({
                        "Valor do Documento (R$)": "R$ {:,.2f}",
                        "Valor Até o Vencimento (R$)": "R$ {:,.2f}",
                        "Desconto (R$)": "R$ {:,.2f}",
                    }),
                    use_container_width=True,
                )

                excel_bytes = gerar_excel_boletos(df_boletos)
                nome_arquivo = f"descontos_boletos_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

                st.download_button(
                    label="⬇️ Baixar planilha (.xlsx)",
                    data=excel_bytes, file_name=nome_arquivo,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="btn_baixar_boletos"
                )
        else:
            st.info("Aguardando o envio de um arquivo PDF.")

if __name__ == "__main__":
    main()